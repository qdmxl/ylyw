#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 docx 的封面区（目录之前的所有段落）替换为一整页封面图片。"""
import sys, os
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX = sys.argv[1]
IMG = sys.argv[2]

# A4 页宽 7560310 EMU = 8.27in；缩放封面图占满正文宽
PAGE_W = 7560310
IMG_W = PAGE_W - 1177290 * 2 + 100000  # 略超出，视觉满版
DPI_W, DPI_H = 1240, 1754
IMG_H = int(IMG_W * DPI_H / DPI_W)

d = Document(DOCX)

# 找“目录”段（Normal 且文本含"目")
toc_idx = None
for i, p in enumerate(d.paragraphs):
    if p.style.name == 'Normal' and '目' in p.text and len(p.text.strip()) <= 6:
        toc_idx = i
        break
print('目录段索引:', toc_idx)
if toc_idx is None:
    raise SystemExit('未找到目录段')

# 删除封面区段落 [0, toc_idx)
for p in d.paragraphs[:toc_idx]:
    p._p.getparent().remove(p._p)

# 在目录段之前插入封面段
toc_el = d.paragraphs[0]._p  # 现在目录段是第 0 段
cover_p = toc_el.makeelement(qn('w:p'), {})
toc_el.addprevious(cover_p)
from docx.text.paragraph import Paragraph
para = Paragraph(cover_p, d.paragraphs[0]._parent)
para.style = d.styles['#M-正文-通用']
para.alignment = WD_ALIGN_PARAGRAPH.CENTER
para.paragraph_format.space_before = 0
para.paragraph_format.space_after = 0
run = para.add_run()
run.add_picture(IMG, width=Emu(IMG_W), height=Emu(IMG_H))

# 让目录另起一页（封面独占一页）
toc_para = d.paragraphs[1]
toc_pPr = toc_para._p.get_or_add_pPr()
if toc_pPr.find(qn('w:pageBreakBefore')) is None:
    toc_pPr.append(toc_pPr.makeelement(qn('w:pageBreakBefore'), {}))

d.save(DOCX)
print('封面已替换为图片:', IMG, '尺寸(EMU):', IMG_W, IMG_H)
