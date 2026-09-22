#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将第10章(汉语语义理解引擎)并入专著，并完成章节重编号。
插入位置：原第9章之后、原第10章(通向通用智能)之前。
新第10章 = 汉语语义理解引擎；原第10章(通向通用智能)→第11章；原第11章(局限)→第12章。
输入:  易理研物-20260910-V3-含量子易理章.docx  +  第10章_汉语语义理解引擎/第10章_汉语语义理解引擎.md
输出:  易理研物-20260911-V4-含汉语语义章.docx
"""
import re, os, sys, copy
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import omml as OMML

BASE = '/home/lijinhan/MXL/科研/ylyw/paper'
CHDIR = os.path.join(BASE, '第10章_汉语语义理解引擎')
SRC = os.path.join(BASE, '易理研物-20260910-V3-含量子易理章.docx')
OUT = os.path.join(BASE, '易理研物-20260911-V4-含汉语语义章.docx')

FONT_SONG = '宋体'
FONT_HEI = '黑体'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def _w(tag):
    return '{%s}%s' % (W_NS, tag)


def set_cn(run, font, size, bold=False, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    rf.set(qn('w:eastAsia'), font)


def _add_inline(p, text, size=11, font=None):
    font = font or FONT_SONG
    tokens = re.split(r'(\*\*[^*]+\*\*|\$[^$]+\$)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            inner = tok[2:-2]
            for sub in re.split(r'(\$[^$]+\$)', inner):
                if not sub:
                    continue
                if sub.startswith('$') and sub.endswith('$') and len(sub) > 2:
                    try:
                        om = OMML.latex_to_omml(sub[1:-1]); p._p.append(om)
                    except Exception:
                        r = p.add_run(sub); set_cn(r, font, size, bold=True)
                else:
                    r = p.add_run(sub); set_cn(r, font, size, bold=True)
        elif tok.startswith('$') and tok.endswith('$') and len(tok) > 2:
            try:
                om = OMML.latex_to_omml(tok[1:-1]); p._p.append(om)
            except Exception:
                r = p.add_run(tok.strip('$')); set_cn(r, font, size)
        else:
            r = p.add_run(tok); set_cn(r, font, size)


def parse_md(path):
    lines = open(path, encoding='utf-8').read().split('\n')
    blocks = []
    i = 0
    while i < len(lines):
        st = lines[i].strip()
        if not st:
            i += 1; continue
        if st == '---':
            i += 1; continue
        m = re.match(r'^(#{1,6})\s+(.*)$', st)
        if m:
            blocks.append(('h', len(m.group(1)), m.group(2).strip())); i += 1; continue
        m = re.match(r'^!\[[^\]]*\]\(([^)]+)\)$', st)
        if m:
            blocks.append(('img', m.group(1), None)); i += 1; continue
        if st.startswith('$$') and st.endswith('$$') and len(st) > 4:
            blocks.append(('formula', st.strip('$').strip(), None)); i += 1; continue
        if st.startswith('>'):
            blocks.append(('quote', st.lstrip('> ').strip(), None)); i += 1; continue
        if st.startswith('|'):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl.append(lines[i].strip()); i += 1
            blocks.append(('table', tbl, None)); continue
        blocks.append(('p', st, None)); i += 1
    return blocks


def append_block(doc, block):
    kind = block[0]
    if kind == 'h':
        lvl = block[1]
        p = doc.add_paragraph()
        p.style = doc.styles['Heading %d' % min(lvl, 3)]
        r = p.add_run(block[2])
        if lvl == 1:
            set_cn(r, FONT_HEI, 18, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif lvl == 2:
            set_cn(r, FONT_HEI, 15, bold=True)
        else:
            set_cn(r, FONT_HEI, 13, bold=True)
        return p
    elif kind == 'img':
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(os.path.join(CHDIR, block[1]), width=Cm(14))
        return p
    elif kind == 'formula':
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            om = OMML.latex_to_omml(block[1]); p._p.append(om)
        except Exception:
            r = p.add_run(block[1]); set_cn(r, FONT_SONG, 11)
        return p
    elif kind == 'quote':
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_inline(p, block[1], 11)
        return p
    elif kind == 'table':
        return append_table(doc, block[1])
    else:
        txt = block[1]
        p = doc.add_paragraph()
        pf = p.paragraph_format
        is_caption = re.match(r'^\*\*(图|表)\d', txt)
        if is_caption:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Pt(0)
            _add_inline(p, txt, 10.5)
        else:
            pf.first_line_indent = Pt(22)
            _add_inline(p, txt, 11)
        pf.space_after = Pt(6); pf.line_spacing = 1.5
        return p


def append_table(doc, rows):
    cells = []
    for r in rows:
        parts = [c.strip() for c in r.strip('|').split('|')]
        if all(re.match(r'^:?-{2,}:?$', c) or c == '' for c in parts):
            continue
        cells.append(parts)
    if not cells:
        return None
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = 1
    for ri, r in enumerate(cells):
        for ci in range(ncol):
            val = r[ci] if ci < len(r) else ''
            cell = t.cell(ri, ci)
            cell.text = ''
            para = cell.paragraphs[0]
            _add_inline(para, val, 10)
            if ri == 0:
                for run in para.runs:
                    run.font.bold = True
    doc.add_paragraph()
    return t


def _set_heading_text(p, new_text):
    runs = p.runs
    if not runs:
        p.add_run(new_text); return
    runs[0].text = new_text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def _rewrite_paragraph_text(p, new_full):
    runs = p.runs
    if not runs:
        p.add_run(new_full); return
    runs[0].text = new_full
    for r in runs[1:]:
        r.text = ''


# ---------- 重编号：原 10→11, 11→12 ----------
def renumber_tblfig(doc):
    """重编号旧章的表/图题：原第10章(表10.x)→11.x；原第11章(表11.x)→12.x。
    逻辑：进入新第10章区域时按序号区分；离开后遭遇的“表10.”实为旧第10章遗留，应→11.x。
    为鲁棒，按当前所在章（H1）判定：若 H1 为第11章（通向通用智能），则表10.→表11.；
    若 H1 为第12章（局限），则表11.→表12.。"""
    cur = None
    for p in doc.paragraphs:
        if p.style.name.startswith('Heading 1'):
            cur = p.text.strip()
            continue
        if cur is None:
            continue
        t = p.text.strip()
        m = re.match(r'^(表|图)(\d+)\.(\d+)(.*)$', t)
        if not m:
            continue
        kind, major, minor, rest = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        new_major = None
        if cur.startswith('第11章') and major == 10:
            new_major = 11
        elif cur.startswith('第12章') and major == 11:
            new_major = 12
        if new_major is None:
            continue
        _rewrite_paragraph_text(p, '%s%d.%s%s' % (kind, new_major, minor, rest))


def renumber(doc):
    in_new_ch10 = False
    for p in doc.paragraphs:
        if not p.style.name.startswith('Heading'):
            continue
        t = p.text.strip()
        m = re.match(r'^第(\d+)章(.*)$', t)
        if m:
            n = int(m.group(1))
            if t.startswith('第10章 汉语语义理解引擎'):
                in_new_ch10 = True
                continue
            if n >= 10:
                in_new_ch10 = False
                _set_heading_text(p, '第%d章%s' % (n + 1, m.group(2)))
            continue
        m2 = re.match(r'^(\d+)\.(\d+)(.*)$', t)
        if m2:
            major = int(m2.group(1))
            if in_new_ch10:
                continue
            if major >= 10:
                _set_heading_text(p, '%d.%s%s' % (major + 1, m2.group(2), m2.group(3)))
            continue


def renumber_intext_refs(doc):
    """正文交叉引用：第10章→第11章、第11章→第12章（跳过新第10章区域）。"""
    in_new = False
    hit = 0
    for p in doc.paragraphs:
        if p.style.name.startswith('Heading'):
            t = p.text.strip()
            if t.startswith('第10章 汉语语义理解引擎'):
                in_new = True; continue
            if t.startswith('第11章') or t.startswith('第12章'):
                in_new = False
        full = p.text
        if not full or in_new:
            continue
        if p.style.name.startswith('Heading'):
            continue
        if re.match(r'^第\d+章：', full.strip()):  # 0.3 导航行，另行处理
            continue
        if '第11章' not in full and '第10章' not in full:
            continue
        new_full = full.replace('第11章', '@@12@@').replace('第10章', '第11章').replace('@@12@@', '第12章')
        if new_full != full:
            _rewrite_paragraph_text(p, new_full)
            hit += 1
    return hit


def update_nav(doc):
    """更新 0.3 结构导航：插入新第10章条目，顺延旧第10/11章。"""
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('第10章：未来——'):
            _rewrite_paragraph_text(p, t.replace('第10章：', '第11章：', 1))
        elif t.startswith('第11章：界——'):
            _rewrite_paragraph_text(p, t.replace('第11章：', '第12章：', 1))
        elif t.startswith('第9章：辩——'):
            new_line = ('第10章：语——汉语语义理解引擎。 从“观物取象”的认知同构出发，'
                        '给出字形-字义-词句-篇章四层嵌套语言架构；以《论语》零训练验证、'
                        '双通道自适应语义底座与嵌套自增长，证明汉语可以在机器中生长。')
            _insert_paragraph_after(p, new_line)


def _insert_paragraph_after(p, text):
    """在段落 p 之后插入一个新段落（拷贝 p 的样式与格式）。"""
    from docx.text.paragraph import Paragraph
    new_p = copy.deepcopy(p._p)
    for t in new_p.findall('.//' + _w('t')):
        t.text = ''
    p._p.addnext(new_p)
    np = Paragraph(new_p, p._parent)
    if np.runs:
        np.runs[0].text = text
        for r in np.runs[1:]:
            r.text = ''
    else:
        np.add_run(text)
    return np


def set_update_fields(path):
    import zipfile, shutil
    tmpf = path + '.tmp'
    with zipfile.ZipFile(path, 'r') as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    s = data['word/settings.xml'].decode('utf-8')
    if 'updateFields' not in s:
        pos = s.find('>', s.find('<w:settings')) + 1
        s = s[:pos] + '<w:updateFields w:val="true"/>' + s[pos:]
        data['word/settings.xml'] = s.encode('utf-8')
    with zipfile.ZipFile(tmpf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmpf, path)


def main():
    doc = Document(SRC)
    body = doc.element.body

    anchor = None
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1' and p.text.strip().startswith('第10章'):
            anchor = p._p
            break
    assert anchor is not None, '未找到第10章锚点'

    blocks = parse_md(os.path.join(CHDIR, '第10章_汉语语义理解引擎.md'))
    before = list(body)
    for b in blocks:
        append_block(doc, b)
    after = list(body)
    before_set = set(id(x) for x in before)
    new_nodes = [x for x in after if id(x) not in before_set]
    for node in new_nodes:
        anchor.addprevious(node)

    renumber(doc)
    renumber_tblfig(doc)
    update_nav(doc)
    hit = renumber_intext_refs(doc)
    print('重编号交叉引用段落数:', hit)

    doc.save(OUT)
    set_update_fields(OUT)
    print('saved:', OUT)
    print('段落:', len(doc.paragraphs), '表:', len(doc.tables), '图:', len(doc.inline_shapes))


if __name__ == '__main__':
    main()
