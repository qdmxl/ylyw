#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
为正式排版文档生成静态目录（带页码 + 点前导符），并保证每章另起一页。
工作流：读取 reformat_v4.py 产出的 docx（含封面 + 正文，无目录），
       在封面之后、正文（序）之前插入目录；可重复运行（幂等）。
两遍法：
  1. 清掉旧目录 -> 渲染 PDF 读取各级标题页码。
  2. 生成目录段落插入，再渲染一次修正页码，最终保存。
用法: python3 finalize_toc.py <docx>
"""
import re, sys, os, subprocess
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = '/home/lijinhan/MXL/科研/ylyw/paper'
DOCX = sys.argv[1] if len(sys.argv) > 1 else BASE + '/易理研物-正式排版-20260911.docx'

MARK = 'TOC_GEN'   # 标记生成的目录段落，便于幂等删除


def render_pdf(docx, outdir):
    os.makedirs(outdir, exist_ok=True)
    for f in os.listdir(outdir):
        try:
            os.remove(os.path.join(outdir, f))
        except OSError:
            pass
    subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf',
                    '--outdir', outdir, docx], check=True,
                   capture_output=True, timeout=900)
    pdfs = [f for f in os.listdir(outdir) if f.endswith('.pdf')]
    return os.path.join(outdir, pdfs[0])


def page_texts(pdf):
    n = int(subprocess.run(['pdfinfo', pdf], capture_output=True, text=True)
            .stdout.split('Pages:')[1].split()[0])
    out = []
    for i in range(1, n + 1):
        t = subprocess.run(['pdftotext', '-f', str(i), '-l', str(i), pdf, '-'],
                           capture_output=True, text=True).stdout
        out.append(t)
    return out


def collect_headings(doc):
    """目录只收录章（Heading 1）与节（Heading 2），三级标题不入目录。"""
    hs = []
    for p in doc.paragraphs:
        s = p.style.name
        t = p.text.strip()
        if not t:
            continue
        if s == 'Heading 1':
            hs.append((1, t))
        elif s == 'Heading 2':
            hs.append((2, t))
    return hs


def locate_pages(heads, pages, start_key=None):
    """在 pages 中按顺序定位每个标题页码（1-based）。
    start_key: 正文起点特征串；从包含它的页开始搜索，避开目录页。"""
    norm = lambda s: re.sub(r'\s+', '', s)
    pn = [norm(t) for t in pages]
    cursor = 0
    if start_key:
        sk = norm(start_key)
        for i, t in enumerate(pn):
            if sk in t:
                cursor = i
                break
    res = []
    for lvl, txt in heads:
        key = norm(txt)
        found = None
        for i in range(cursor, len(pn)):
            if key and key in pn[i]:
                found = i + 1
                cursor = i
                break
        res.append((lvl, txt, found))
    return res


def body_anchor(doc):
    """返回第一个正文标题（Heading 4 或 Heading 1）的 XML 元素。"""
    for p in doc.paragraphs:
        if p.style.name in ('Heading 4', 'Heading 1'):
            return p._p
    raise RuntimeError('未找到正文锚点')


def remove_old_toc(doc):
    """删除所有带 MARK 书签/标记的目录段落。"""
    body = doc.element.body
    for p in list(doc.paragraphs):
        if p._p.find('.//' + qn('w:bookmarkStart')) is not None:
            for bs in p._p.findall('.//' + qn('w:bookmarkStart')):
                if bs.get(qn('w:name'), '').startswith(MARK):
                    body.remove(p._p)
                    break
        # 目录标题（样式）也删
        if p.style.name == '#P-目录-标题' and p.text.strip().startswith('目'):
            # 仅当它是我们生成的分页区附近；稳妥起见也删
            if p._p.find('.//' + qn('w:bookmarkStart')) is None:
                # 判断是否已被上面删除
                if p._p.getparent() is not None:
                    body.remove(p._p)


def new_para(doc, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    doc.element.body.remove(p._p)
    return p


def build_toc(doc, entries):
    """生成真正的目录域（TOC field）：
    结构 = 域开始(begin) + 域指令(instrText) + 分隔(separate)
           + 缓存结果（静态条目，可直接显示） + 域结束(end)。
    这样目录可自动生成/可更新（Word 中 F9 刷新），且打开即见。"""
    anchor = body_anchor(doc)

    nodes = []

    # 目录标题
    title = new_para(doc, '#P-目录-标题')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run('目　录')
    nodes.append(title._p)

    # 域开始
    p_begin = new_para(doc)
    rb = p_begin.add_run()
    fld_b = OxmlElement('w:fldChar')
    fld_b.set(qn('w:fldCharType'), 'begin')
    rb._r.append(fld_b)
    nodes.append(p_begin._p)

    # 域指令
    p_instr = new_para(doc)
    ri = p_instr.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = r' TOC \o "1-2" \h \z \u '
    ri._r.append(it)
    nodes.append(p_instr._p)

    # 分隔
    p_sep = new_para(doc)
    rs = p_sep.add_run()
    fld_s = OxmlElement('w:fldChar')
    fld_s.set(qn('w:fldCharType'), 'separate')
    rs._r.append(fld_s)
    nodes.append(p_sep._p)

    # 缓存结果：静态条目（打开即显示，F9 可刷新）
    for lvl, txt, pg in entries:
        p = new_para(doc)
        p.paragraph_format.left_indent = Pt(12 * (lvl - 1))
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(14.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(txt)
        r.font.size = Pt(12 if lvl == 1 else 10.5)
        r.bold = (lvl == 1)
        if pg:
            p.add_run('\t' + str(pg))
        nodes.append(p._p)

    # 域结束
    p_end = new_para(doc)
    re_ = p_end.add_run()
    fld_e = OxmlElement('w:fldChar')
    fld_e.set(qn('w:fldCharType'), 'end')
    re_._r.append(fld_e)
    nodes.append(p_end._p)

    # 目录后分页
    brk = new_para(doc)
    brk.add_run().add_break(WD_BREAK.PAGE)
    nodes.append(brk._p)

    # 打标记（幂等性）
    for n in nodes:
        bs = OxmlElement('w:bookmarkStart')
        bs.set(qn('w:id'), '9000')
        bs.set(qn('w:name'), MARK)
        n.insert(0, bs)
        be = OxmlElement('w:bookmarkEnd')
        be.set(qn('w:id'), '9000')
        n.append(be)

    for n in nodes:
        anchor.addprevious(n)


def add_page_break_before_chapters(doc):
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1':
            pPr = p._p.get_or_add_pPr()
            if pPr.find(qn('w:pageBreakBefore')) is None:
                pPr.append(OxmlElement('w:pageBreakBefore'))


def main():
    START_KEY = '从此跟AI一起'  # 序正文首句，作为正文起点
    d = Document(DOCX)
    add_page_break_before_chapters(d)
    remove_old_toc(d)
    d.save(DOCX)

    heads = collect_headings(Document(DOCX))
    print('待排目录标题:', len(heads))
    pdf = render_pdf(DOCX, '/tmp/toc_r1')
    located = locate_pages(heads, page_texts(pdf), start_key=START_KEY)

    d2 = Document(DOCX)
    build_toc(d2, located)
    d2.save(DOCX)
    # 第二轮修正页码（目录占页导致偏移）
    pdf2 = render_pdf(DOCX, '/tmp/toc_r2')
    located2 = locate_pages(heads, page_texts(pdf2), start_key=START_KEY)
    d3 = Document(DOCX)
    remove_old_toc(d3)
    build_toc(d3, located2)
    d3.save(DOCX)
    print('目录完成。前若干条：')
    for it in located2[:12]:
        print('  ' * (it[0] - 1) + f'[{it[0]}] {it[1]} -> {it[2]}')
    print('...')
    for it in located2[-6:]:
        print('  ' * (it[0] - 1) + f'[{it[0]}] {it[1]} -> {it[2]}')


if __name__ == '__main__':
    main()
